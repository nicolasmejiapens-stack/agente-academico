"""Conexiones: Google (Drive y Sheets), Telegram y Claude."""
import io
import json
import logging
import os
import re
import time

import anthropic
import docx
import gspread
import requests
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from pypdf import PdfReader

log = logging.getLogger("gestor")
SCOPES = ["https://www.googleapis.com/auth/drive", "https://www.googleapis.com/auth/spreadsheets"]
GDOC = "application/vnd.google-apps.document"


class ErrorIA(Exception):
    pass


# ---------- Google ----------
def _creds():
    return Credentials.from_service_account_info(json.loads(os.environ["GOOGLE_SA_JSON"]), scopes=SCOPES)


def abrir_sheet():
    return gspread.authorize(_creds()).open_by_key(os.environ["SHEET_ID"])


def drive():
    return build("drive", "v3", credentials=_creds(), cache_discovery=False)


def leer_config(sh):
    filas = sh.worksheet("Config").get_all_values(value_render_option="UNFORMATTED_VALUE")[1:]
    cfg = {}
    for f in filas:
        if len(f) >= 2 and str(f[0]).strip():
            v = f[1]
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            cfg[str(f[0]).strip()] = str(v).strip()
    return cfg


def hoja_vistos(sh):
    try:
        return sh.worksheet("Archivos_Vistos")
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title="Archivos_Vistos", rows=500, cols=7)
        ws.update(range_name="A1:G1", values=[["ID_Archivo", "Materia", "Nombre", "Tipo", "Estado", "Intentos", "Fecha"]])
        return ws


def agregar(sh, pestana, filas):
    if filas:
        sh.worksheet(pestana).append_rows(filas, value_input_option="USER_ENTERED")


def listar_archivos(svc, carpeta_id, profundidad=2):
    out, token = [], None
    q = f"'{carpeta_id}' in parents and trashed=false"
    while True:
        r = svc.files().list(q=q, fields="nextPageToken, files(id,name,mimeType,size)", pageSize=200,
                             pageToken=token, supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
        for f in r.get("files", []):
            if f["mimeType"] == "application/vnd.google-apps.folder":
                if profundidad > 1:
                    out += listar_archivos(svc, f["id"], profundidad - 1)
            else:
                out.append(f)
        token = r.get("nextPageToken")
        if not token:
            return out


def soportado(f):
    n = f["name"].lower()
    return f["mimeType"] == GDOC or n.endswith(".pdf") or n.endswith(".docx")


def descargar(svc, f):
    if f["mimeType"] == GDOC:
        req = svc.files().export_media(fileId=f["id"], mimeType="text/plain")
    else:
        req = svc.files().get_media(fileId=f["id"], supportsAllDrives=True)
    buf, dl, done = io.BytesIO(), None, False
    dl = MediaIoBaseDownload(buf, req)
    while not done:
        _, done = dl.next_chunk()
    return buf.getvalue()


def extraer_texto(f, datos, max_paginas=None):
    nombre = f["name"].lower()
    if f["mimeType"] == GDOC:
        return datos.decode("utf-8", "ignore")
    if nombre.endswith(".pdf"):
        pags = PdfReader(io.BytesIO(datos)).pages
        pags = list(pags)[:max_paginas] if max_paginas else list(pags)
        return "\n".join((p.extract_text() or "") for p in pags)
    if nombre.endswith(".docx"):
        d = docx.Document(io.BytesIO(datos))
        partes = [p.text for p in d.paragraphs]
        for t in d.tables:
            for fila in t.rows:
                partes.append(" | ".join(c.text.strip() for c in fila.cells))
        return "\n".join(partes)
    return ""


# ---------- Telegram ----------
def telegram(cfg, texto):
    try:
        r = requests.post(
            f"https://api.telegram.org/bot{os.environ['TELEGRAM_TOKEN']}/sendMessage",
            json={"chat_id": cfg.get("telegram_chat_id", ""), "text": texto[:4000],
                  "disable_web_page_preview": True},
            timeout=20)
        if not r.ok:
            log.error("Telegram respondió %s", r.status_code)
        return r.ok
    except Exception as e:
        log.error("Telegram falló: %s", type(e).__name__)
        return False


def tg(metodo, **datos):
    """Llamada genérica a la API de Telegram. Nunca registra el token."""
    try:
        r = requests.post(f"https://api.telegram.org/bot{os.environ['TELEGRAM_TOKEN']}/{metodo}",
                          json=datos, timeout=30)
        return r.json() if r.content else {}
    except Exception as e:
        log.error("Telegram %s falló: %s", metodo, type(e).__name__)
        return {}


def tg_documento(chat_id, nombre, datos, caption="", reply_markup=None):
    """Envía un archivo a Telegram. Devuelve el file_id o None. Límite de Telegram: 50 MB."""
    try:
        limpio = re.sub(r"[^\w\-. ]+", "", nombre).strip() or "documento.pdf"
        form = {"chat_id": chat_id, "caption": caption[:1000]}
        if reply_markup:
            form["reply_markup"] = json.dumps(reply_markup)
        r = requests.post(f"https://api.telegram.org/bot{os.environ['TELEGRAM_TOKEN']}/sendDocument",
                          data=form, files={"document": (limpio, datos)}, timeout=240)
        j = r.json()
        if j.get("ok"):
            return j["result"]["document"]["file_id"]
        log.error("sendDocument respondió %s", r.status_code)
    except Exception as e:
        log.error("sendDocument falló: %s", type(e).__name__)
    return None


# ---------- IA con reintentos y respaldo ----------
def _llamar(modelo, system, prompt, max_tokens, intentos=4):
    cliente = anthropic.Anthropic(max_retries=0, timeout=120)
    espera = 3
    ultimo = None
    for _ in range(intentos):
        try:
            r = cliente.messages.create(model=modelo, max_tokens=max_tokens, system=system,
                                        messages=[{"role": "user", "content": prompt}])
            return "".join(b.text for b in r.content if b.type == "text")
        except anthropic.APIStatusError as e:
            ultimo = e
            if not (e.status_code in (408, 409, 429) or e.status_code >= 500):
                raise
        except anthropic.APIConnectionError as e:
            ultimo = e
        time.sleep(espera)
        espera *= 2
    raise ErrorIA(f"La IA no respondió tras {intentos} intentos ({type(ultimo).__name__})")


def llamar_ia(cfg, system, prompt, max_tokens=1000):
    principal = cfg.get("modelo_principal", "claude-haiku-4-5-20251001")
    respaldo = cfg.get("modelo_respaldo", "")
    try:
        return _llamar(principal, system, prompt, max_tokens)
    except anthropic.AuthenticationError:
        raise ErrorIA("API key de Anthropic inválida o sin permisos")
    except Exception as e:
        log.warning("Modelo principal falló (%s); probando respaldo", type(e).__name__)
        if not respaldo:
            raise ErrorIA("Falló el modelo principal y no hay respaldo")
        try:
            return _llamar(respaldo, system, prompt, max_tokens)
        except Exception as e2:
            raise ErrorIA(f"Fallaron principal y respaldo ({type(e2).__name__})")


def pedir_json(cfg, system, prompt, max_tokens=1000):
    txt = llamar_ia(cfg, system, prompt, max_tokens).strip()
    txt = re.sub(r"^```(?:json)?|```$", "", txt, flags=re.M).strip()
    i, j = txt.find("{"), txt.rfind("}")
    try:
        return json.loads(txt[i:j + 1])
    except Exception:
        raise ErrorIA("La IA devolvió una respuesta que no es JSON válido")


SYS_JSON = ("Respondes SOLO con JSON válido, sin texto adicional ni bloques de código. "
            "Si un dato no aparece en el texto, usa null. Nunca inventes datos.")


def clasificar(cfg, nombre, muestra):
    prompt = (f"Clasifica este archivo de una materia universitaria de psicología.\n"
              f"Nombre: {nombre}\nInicio del contenido:\n{muestra}\n\n"
              'Devuelve {"tipo": "programa" | "lectura" | "otro"}. '
              '"programa" = syllabus o programa oficial de la materia (cronograma, evaluaciones, bibliografía). '
              '"lectura" = artículo, capítulo o libro para leer. "otro" = cualquier otra cosa.')
    t = pedir_json(cfg, SYS_JSON, prompt, 100).get("tipo")
    return t if t in ("programa", "lectura", "otro") else "lectura"


def extraer_programa(cfg, materia, inicio_semestre, texto):
    prompt = f"""Extrae la estructura del programa (syllabus) de la materia "{materia}".
El semestre empieza el {inicio_semestre}. Úsalo para deducir el año de las fechas.

Devuelve exactamente este JSON:
{{"semanas": [{{"semana": 1, "unidad": "...", "tema": "...",
   "lecturas": [{{"autor": "...", "titulo": "...", "paginas_impresas": "45-72" o null, "obligatoria": true}}]}}],
 "evaluaciones": [{{"nombre": "...", "tipo": "parcial" | "quiz" | "entrega", "semana": 8 o null,
   "fecha": "AAAA-MM-DD" o null, "porcentaje": 20 o null}}]}}

Reglas:
- Máximo 16 semanas, numeradas desde 1.
- "obligatoria" es true salvo que el texto diga explícitamente complementaria, sugerida u opcional.
- "fecha" solo si el programa da una fecha explícita; si solo dice la semana, deja fecha en null.
- "porcentaje" es un número entre 0 y 100, sin el signo %.
- Examen final y parciales son "parcial"; talleres, ensayos y trabajos son "entrega".

TEXTO DEL PROGRAMA:
{texto[:60000]}"""
    return pedir_json(cfg, SYS_JSON, prompt, 7000)


SYS_RES = ("Eres un asistente académico especializado en psicología. Escribes en español claro y riguroso. "
           "Usas solo información presente en el texto; si algo no aparece, lo dices. "
           "No uses formato Markdown (nada de asteriscos ni almohadillas).")


def asociar_lectura(cfg, nombre, muestra, candidatas):
    """Devuelve la lista de IDs del programa que corresponden al archivo (puede ser más de una por libro)."""
    lista = "\n".join(f"{r['id']} | semana {r['semana']} | {r['autor']} | {r['titulo']} | "
                      f"pp. {r['pi'] if r['pi'] is not None else '?'}-{r['pf'] if r['pf'] is not None else '?'}"
                      for r in candidatas[:80])
    prompt = (f'Un estudiante subió el archivo "{nombre}". Inicio de su contenido:\n{muestra[:1500]}\n\n'
              f"Lecturas del programa que aún no tienen archivo:\n{lista}\n\n"
              'Devuelve {"ids": ["ID exacto", ...], "confianza": "alta" | "media" | "baja"}. '
              "Incluye todas las lecturas que salen de ESTE archivo (un libro puede cubrir varios capítulos de varias semanas). "
              "Si ninguna corresponde con claridad, devuelve ids vacío.")
    j = pedir_json(cfg, SYS_JSON, prompt, 300)
    validos = {r["id"] for r in candidatas}
    if j.get("confianza") not in ("alta", "media"):
        return []
    return [i for i in (j.get("ids") or []) if i in validos]


def resumir(cfg, materia, semana, titulo, autor, texto, profundo=True):
    import logica as L
    ctx = f'Materia: {materia}. Semana {semana}. Lectura: "{titulo}" de {autor}.'
    if len(texto) > 90000:
        bloques = L.partir(texto, 60000)
        parciales = []
        for i, b in enumerate(bloques, 1):
            parciales.append(llamar_ia(
                cfg, SYS_RES,
                f"{ctx}\nResume este bloque ({i}/{len(bloques)}) conservando conceptos clave, autores, teorías, "
                f"hallazgos y ejemplos importantes. Máximo 700 palabras.\n\nTEXTO:\n{b}", 1500))
        texto = "\n\n".join(parciales)
    if profundo:
        prompt = (f"{ctx}\nEscribe un resumen académico profundo con estas secciones, cada una con su título en MAYÚSCULAS:\n"
                  "IDEA CENTRAL\nCONCEPTOS CLAVE (definición breve de cada uno)\nAUTORES Y TEORÍAS\n"
                  "HALLAZGOS O EVIDENCIA\nCRÍTICAS, LÍMITES Y RELACIÓN CON OTROS TEMAS\n"
                  "PREGUNTAS PROBABLES DE PARCIAL (3 a 5)\n"
                  'Si una sección no aplica, escribe "No aparece en este fragmento".\n\nTEXTO:\n' + texto)
        return llamar_ia(cfg, SYS_RES, prompt, 3500).strip()
    prompt = (f"{ctx}\nEs una lectura sugerida (no obligatoria). Resume en máximo 200 palabras: idea central, "
              "3 a 5 conceptos clave y autores mencionados.\n\nTEXTO:\n" + texto)
    return llamar_ia(cfg, SYS_RES, prompt, 900).strip()
